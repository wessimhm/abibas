#!/usr/bin/env python3
"""
Contrôle déterministe d'une fiche de connaissance convertie depuis un Word ou un Excel.

Usage :
    python3 .claude/skills/ajouter-connaissance/scripts/verifier_conversion.py "travail/a-ajouter/3-processus/PR-LAB-04 - Etalonnage des instruments.md"
    (le document d'origine est lu dans la ligne <!-- source : ... --> de la fiche,
     ou donné en second argument)

Contrôles :
- chaque nombre du document d'origine est dans la fiche (aucune valeur perdue) ;
- aucun nombre de la fiche n'est absent du document (aucune valeur inventée ou convertie) ;
- virgule décimale, en-tête ref et source, mention « Brouillon » ;
- chaque image du document a son [À COMPLÉTER] ; commentaires et modifications suivies signalés.
On compte les occurrences : une valeur qui revient deux fois dans le document doit revenir deux fois.
Ne comptent pas : les commentaires HTML, la ligne « Brouillon », les numéros de liste (« 1. »),
et ce qui est entre crochets [À COMPLÉTER …].
Code de sortie : 0 si tout passe, 1 sinon. Lecture seule : ce script ne modifie rien.
"""
import datetime
import os
import re
import sys
from collections import Counter

for _flux in (sys.stdout, sys.stderr):  # Windows : affichage UTF-8 (−, ≥, Ω)
    try:
        _flux.reconfigure(encoding="utf-8")
    except Exception:
        pass


def racine_depot():
    d = os.path.dirname(os.path.abspath(__file__))
    while d != os.path.dirname(d):
        if os.path.isdir(os.path.join(d, "originaux")):
            return d
        d = os.path.dirname(d)
    return os.getcwd()


RACINE = racine_depot()
NOMBRE = re.compile(r"\d+(?: \d{3})*(?:,\d+)?")


def normaliser(t):
    return t.replace(" ", " ").replace(" ", " ").replace(" ", " ")


def nombres(t):
    # « 1 200 » et « 1200 » sont le même nombre (espace des milliers du glossaire)
    return Counter(n.replace(" ", "") for n in NOMBRE.findall(normaliser(t)))


def val(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return (str(int(v)) if v.is_integer() else str(v)).replace(".", ",")
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d") if not (v.hour or v.minute) else v.strftime("%Y-%m-%d %H:%M")
    return str(v)


def lire_source(chemin):
    """Renvoie (texte, nb_images, nb_commentaires, nb_modifs_suivies)."""
    if chemin.lower().endswith(".docx"):
        import docx
        from docx.oxml.ns import qn
        d = docx.Document(chemin)
        morceaux, vus, gardes = [], set(), []  # gardes : garde les cellules en vie (id stable)
        for el in d.element.body.iterchildren():
            tag = el.tag.split("}")[-1]
            if tag == "p":
                morceaux.append("".join(t.text or "" for t in el.iter(qn("w:t"))))
            elif tag == "tbl":
                from docx.table import Table
                for r in Table(el, d).rows:
                    for c in r.cells:  # une cellule fusionnée ne compte qu'une fois
                        if id(c._tc) in vus:
                            continue
                        vus.add(id(c._tc))
                        gardes.append(c._tc)
                        morceaux.append(c.text)
        corps = d.element.body
        images = len(corps.findall(".//" + qn("w:drawing"))) + len(corps.findall(".//" + qn("w:pict")))
        modifs = len(corps.findall(".//" + qn("w:ins"))) + len(corps.findall(".//" + qn("w:del")))
        commentaires = 0
        for rel in d.part.rels.values():
            if rel.reltype.endswith("/comments"):
                commentaires = len(rel.target_part.blob.decode("utf-8", "ignore").split("<w:comment ")) - 1
        return "\n".join(morceaux), images, commentaires, modifs
    if chemin.lower().endswith(".xlsx"):
        import openpyxl
        lignes = []
        for ws in openpyxl.load_workbook(chemin, data_only=True):
            for r in ws.iter_rows(values_only=True):
                lignes.append(" | ".join(val(v) for v in r))
        return "\n".join(lignes), 0, 0, 0
    raise ValueError("format non pris en charge (attendu : .docx ou .xlsx)")


def texte_fiche_a_compter(t):
    t = re.sub(r"<!--.*?-->", "", t, flags=re.S)
    t = re.sub(r"\[À COMPLÉTER[^\]]*\]", "", t)
    garde = []
    for l in t.splitlines():
        if re.search(r"brouillon rédigé avec l.aide de l.IA", l, re.I):
            continue
        if re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", l) and "-" in l:  # séparateur de tableau
            continue
        l = re.sub(r"^\s*(?:>\s*)?(?:\d+[.)]|[-*+])\s+", "", l)  # numéro ou puce de liste
        l = re.sub(r"^#+\s*", "", l)
        garde.append(l)
    return "\n".join(garde)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    fiche = sys.argv[1]
    if not os.path.isabs(fiche):
        fiche = os.path.join(RACINE, fiche)
    try:
        t = open(fiche, encoding="utf-8").read()
    except OSError as e:
        print(f"ERREUR : fiche illisible ({fiche}) : {e}")
        sys.exit(1)

    m = re.search(r"<!--\s*source\s*:\s*(.+?)\s*-->", t)
    source = sys.argv[2] if len(sys.argv) > 2 else (m.group(1) if m else "")
    ok_total = total = 0

    def check(ok, nom, detail=""):
        nonlocal ok_total, total
        total += 1
        ok_total += bool(ok)
        print(f"PASS  {nom}" if ok else f"FAIL  {nom} : {detail}")

    print(f"Fiche  : {os.path.relpath(fiche, RACINE)}")
    print(f"Source : {source or '(introuvable)'}\n")
    check(re.search(r"<!--\s*ref:", t), "En-tête <!-- ref: … --> présent", "ajouter <!-- ref: <code> rév. <X> --> en ligne 1")
    check(bool(source), "Source indiquée", "ajouter <!-- source : originaux/… --> en ligne 2")
    if not source:
        print(f"\nRÉSULTAT : {ok_total}/{total} contrôles passés, {total - ok_total} en échec")
        sys.exit(1)
    chemin_source = source if os.path.isabs(source) else os.path.join(RACINE, source)
    try:
        texte_src, images, commentaires, modifs = lire_source(chemin_source)
    except Exception as e:
        print(f"ERREUR : document d'origine illisible ({source}) : {e}")
        sys.exit(1)

    n_src = nombres(texte_src)
    compte = texte_fiche_a_compter(t)
    n_md = nombres(compte)
    manquants = n_src - n_md
    en_trop = n_md - n_src
    fmt = lambda c: ", ".join(f"« {k} »" + (f" ×{v}" if v > 1 else "") for k, v in sorted(c.items()))
    check(not manquants, f"Aucun nombre perdu ({sum(n_src.values())} nombres dans le document)",
          "absents de la fiche : " + fmt(manquants))
    check(not en_trop, "Aucun nombre en trop (inventé ou converti)", "absents du document : " + fmt(en_trop))
    points = re.findall(r"\d\.\d", compte)
    check(not points, "Virgule décimale (aucun « 0.8 »)", f"{len(points)} nombre(s) avec un point décimal")
    check(re.search(r"brouillon rédigé avec l.aide de l.IA, à vérifier", t, re.I),
          "Mention « Brouillon rédigé avec l’aide de l’IA, à vérifier »", "mention absente")
    a_completer = len(re.findall(r"\[À COMPLÉTER[^\]]*image", t, re.I))
    check(a_completer >= images, f"Images signalées ({images} dans le document)",
          f"{images - a_completer} image(s) sans [À COMPLÉTER : image non convertie, …]")
    if commentaires or modifs:
        print(f"ATTENTION  le document contient {commentaires} commentaire(s) et {modifs} modification(s) suivie(s) : "
              "l'outil de lecture ne les voit pas, à relire dans Word.")

    print(f"\nRÉSULTAT : {ok_total}/{total} contrôles passés, {total - ok_total} en échec")
    sys.exit(0 if ok_total == total else 1)


if __name__ == "__main__":
    main()
