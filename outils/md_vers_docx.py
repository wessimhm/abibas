#!/usr/bin/env python3
"""Convertit un document Markdown simple en Word, au style des documents Abibas Sport.

Usage : python3 outils/md_vers_docx.py travail/RA-L500-01.md            -> travail/RA-L500-01.docx
        python3 outils/md_vers_docx.py source.md sortie.docx

Prend en charge : # titre, _sous-titre_ (ligne juste après le titre), ## et ### titres,
paragraphes, listes « - » et « 1. », tableaux « | », citations « > », **gras**.
Une première ligne <!-- ref: ... --> donne la référence affichée en en-tête.
Chaque document porte en pied de page la mention « Document fictif de formation ».
"""
import re
import sys
for _flux in (sys.stdout, sys.stderr):  # Windows : affichage UTF-8 (−, ≥, Ω)
    try:
        _flux.reconfigure(encoding="utf-8")
    except Exception:
        pass

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    print("ERREUR : installez la dépendance : pip install python-docx")
    sys.exit(3)

NAVY = RGBColor(0x16, 0x23, 0x3F)
GRIS = RGBColor(0x6B, 0x72, 0x80)
POLICE = "Calibri"
MENTION = ("Document fictif de formation · Abibas Sport, la Ligue de foot de Belmont, "
           "les personnes et toutes les valeurs sont inventés")


def _police(style, taille=None, couleur=None, gras=None):
    style.font.name = POLICE
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), POLICE)
    if taille:
        style.font.size = Pt(taille)
    if couleur is not None:
        style.font.color.rgb = couleur
    if gras is not None:
        style.font.bold = gras


def _ombre(cellule, couleur="E6E9F0"):
    tcpr = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    tcpr.append(shd)


def _texte(par, texte, italique=False, couleur=None, taille=None, gras_tout=False):
    texte = texte.replace("`", "")
    for m in re.split(r"(\*\*[^*]+\*\*)", texte):
        if not m:
            continue
        gras = m.startswith("**") and m.endswith("**") and len(m) > 4
        r = par.add_run(m[2:-2] if gras else m)
        r.bold = gras or gras_tout
        r.italic = italique
        if couleur is not None:
            r.font.color.rgb = couleur
        if taille:
            r.font.size = Pt(taille)


def _cellules(ligne):
    ligne = ligne.strip()
    if ligne.startswith("|"):
        ligne = ligne[1:]
    if ligne.endswith("|"):
        ligne = ligne[:-1]
    return [c.strip() for c in ligne.split("|")]


def _tableau(d, lignes):
    rangs = [_cellules(l) for l in lignes]
    rangs = [r for r in rangs if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    entete = True
    if rangs and not any(rangs[0]):  # en-tête vide : tableau « champ | valeur »
        rangs, entete = rangs[1:], False
    n = max(len(r) for r in rangs)
    t = d.add_table(rows=len(rangs), cols=n)
    t.style = "Table Grid"
    for i, r in enumerate(rangs):
        for j in range(n):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            _texte(p, r[j] if j < len(r) else "", taille=9.5, gras_tout=(i == 0 and entete))
            if i == 0 and entete:
                _ombre(c)
    d.add_paragraph()


def convertir(src, out=None):
    brut = open(src, encoding="utf-8").read().splitlines()
    ref = ""
    if brut and brut[0].strip().startswith("<!--"):
        m = re.search(r"ref:\s*(.*?)\s*-->", brut[0])
        ref = m.group(1) if m else ""
        brut = brut[1:]

    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
    _police(d.styles["Normal"], 10.5)
    _police(d.styles["Heading 1"], 14, NAVY, True)
    _police(d.styles["Heading 2"], 12, NAVY, True)
    _police(d.styles["Heading 3"], 11, NAVY, True)
    for nom in ("List Bullet", "List Number", "Table Grid"):
        try:
            _police(d.styles[nom])
        except KeyError:
            pass

    sec = d.sections[0]
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _texte(hp, "Abibas Sport" + (f" · {ref}" if ref else ""), couleur=GRIS, taille=8.5)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _texte(fp, MENTION, italique=True, couleur=GRIS, taille=8)

    i, titre_vu = 0, False
    while i < len(brut):
        ligne = brut[i]
        s = ligne.strip()
        if not s or s == "_Document fictif de formation._":
            i += 1
            continue
        if s.startswith("|"):
            bloc = []
            while i < len(brut) and brut[i].strip().startswith("|"):
                bloc.append(brut[i])
                i += 1
            _tableau(d, bloc)
            continue
        m = re.match(r"^(#+)\s*(.*)", s)
        if m:
            niveau = len(m.group(1))
            if niveau == 1 and not titre_vu:
                p = d.add_paragraph()
                _texte(p, m.group(2), couleur=NAVY, taille=18, gras_tout=True)
                titre_vu = True
                if i + 1 < len(brut) and re.match(r"^_.*_$", brut[i + 1].strip()):
                    p = d.add_paragraph()
                    _texte(p, brut[i + 1].strip()[1:-1], italique=True, couleur=GRIS, taille=10)
                    i += 1
            else:
                d.add_heading(m.group(2), min(max(niveau - 1, 1), 3))
            i += 1
            continue
        if re.match(r"^[-*]\s+", s):
            _texte(d.add_paragraph(style="List Bullet"), re.sub(r"^[-*]\s+", "", s))
        elif re.match(r"^\d+\.\s+", s):
            _texte(d.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", s))
        elif s.startswith(">"):
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            _texte(p, s.lstrip("> ").strip(), italique=True)
        elif re.match(r"^_.*_$", s):
            _texte(d.add_paragraph(), s[1:-1], italique=True, couleur=GRIS)
        else:
            _texte(d.add_paragraph(), s)
        i += 1

    out = out or re.sub(r"\.md$", ".docx", src)
    d.save(out)
    return out


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    print(f"Créé : {convertir(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)}")
