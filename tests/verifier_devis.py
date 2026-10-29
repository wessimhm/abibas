#!/usr/bin/env python3
"""
Vérificateur déterministe de l'offre « série 2 » (Ligue de foot de Belmont, fictif).

Usage :
    python3 tests/verifier_devis.py                         # vérifie le dernier travail/offre*.md
    python3 tests/verifier_devis.py travail/offre-serie-2.md

Montant attendu (grille 2026) : 1 450 + 1 200 + 1 650 = 4 300 $ ; 2e spécimen 1 505 $ ;
rapport 950 $ ; total 6 755 $ avant taxes ; délai 10 jours ouvrables, pas d'urgence.
Code de sortie : 0 si tout passe, 1 sinon.
"""
import glob
import os
import re
import sys
for _flux in (sys.stdout, sys.stderr):  # Windows : affichage UTF-8 (−, ≥, Ω)
    try:
        _flux.reconfigure(encoding="utf-8")
    except Exception:
        pass

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def lire(chemin):
    if chemin.lower().endswith(".docx"):
        import docx
        d = docx.Document(chemin)
        t = [p.text for p in d.paragraphs]
        for tb in d.tables:
            for r in tb.rows:
                t.append(" | ".join(c.text for c in r.cells))
        brut = "\n".join(t)
    else:
        brut = open(chemin, encoding="utf-8").read()
    return brut.replace(" ", " ").replace(" ", " ").replace("’", "'").replace("**", "")


def montant(texte, n):
    """Le montant n écrit « 6 755 », « 6755 » ou « 6 755,00 »."""
    m = f"{n:,}".replace(",", r"[\s.]?")
    return re.search(rf"(?<![\d,]){m}(?:,00)?\s*\$", texte) is not None


def main():
    if len(sys.argv) > 1:
        cible = sys.argv[1]
    else:
        offres = sorted(glob.glob(os.path.join(RACINE, "travail", "offre*.md")), key=os.path.getmtime)
        if not offres:
            print("ERREUR : aucune offre dans travail/ ; donner le fichier en argument.")
            sys.exit(3)
        cible = offres[-1]
    t = lire(cible)
    ok_total, total = 0, 0
    print(f"Vérification de : {os.path.basename(cible)}\n")

    def check(ok, nom, detail=""):
        nonlocal ok_total, total
        total += 1
        ok_total += bool(ok)
        print(f"PASS  {nom}" if ok else f"FAIL  {nom} : {detail}")

    check(montant(t, 6755) and re.search(r"avant taxes", t, re.I), "Total = 6 755 $ avant taxes",
          "le total attendu (grille 2026) est 6 755 $ avant taxes")
    check(montant(t, 4300), "Sous-total des essais = 4 300 $", "rebond 1 450 + absorption d’eau 1 200 + froid 1 650 = 4 300 $")
    check(montant(t, 1505), "Spécimen supplémentaire = 1 505 $ (35 % de 4 300 $)", "4 300 × 0,35 = 1 505 $")
    check(montant(t, 950), "Rapport et revue = 950 $", "un rapport à 950 $")
    lignes_urgence = [l for l in t.splitlines() if re.search(r"urgence", l, re.I)]
    appliquee = [l.strip() for l in lignes_urgence
                 if not re.search(r"\b(pas|sans|aucune?|non applicable|ne s'applique|n'est pas)\b", l, re.I)]
    check(not appliquee and not montant(t, 8106) and not montant(t, 1351),
          "Aucune majoration d'urgence (délai de 10 jours ouvrables, demande à 15)",
          "majoration d'urgence appliquée à tort : " + (appliquee[0] if appliquee else "8 106 $ ou 1 351 $"))
    check(re.search(r"10 jours ouvrables", t, re.I), "Délai annoncé : 10 jours ouvrables",
          "5 jours d'essais + 5 jours de rapport = 10 jours ouvrables")
    check(re.search(r"taxes en sus", t, re.I) and re.search(r"30 jours", t), "Taxes en sus et validité 30 jours",
          "mentions de la grille manquantes")
    check(re.search(r"rév\.\s*C", t), "Spécification citée en rév. C", "SPEC-LFB-BAL-01 rév. C attendu")
    check(re.search(r"brouillon|à valider", t, re.I) and re.search(r"Karim Haddad", t),
          "Brouillon marqué, validation par le chargé de projet", "l'offre doit être validée et envoyée par Karim Haddad")

    print(f"\nRÉSULTAT : {ok_total}/{total} contrôles passés, {total - ok_total} en échec")
    sys.exit(0 if ok_total == total else 1)


if __name__ == "__main__":
    main()
