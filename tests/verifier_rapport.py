#!/usr/bin/env python3
"""
Vérificateur déterministe d'un rapport d'essai LP-500 (démo formation IA, vendredi 2).

Le modèle rédige et juge ; ce script calcule et vérifie. Même entrée, même résultat.

Usage :
    python3 tests/verifier_rapport.py                         # vérifie le brouillon de essais/
    python3 tests/verifier_rapport.py travail/RA-L500-01.md   (ou .docx)

Sortie : une ligne PASS ou FAIL par contrôle, puis le résultat.
Code de sortie : 0 si tout passe, 1 sinon, 3 en cas d'erreur de fichier.
Dépendances : pip install python-docx openpyxl
Tous les documents sont fictifs (Abibas Sport, Ligue de foot de Belmont).
"""
import glob
import os
import re
import sys
for _flux in (sys.stdout, sys.stderr):  # Windows : affichage UTF-8 (−, ≥, Ω)
    try:
        _flux.reconfigure(encoding="utf-8")
    except Exception:
        pass

try:
    import docx  # python-docx
    import openpyxl
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    print("ERREUR : installez les dépendances : pip install python-docx openpyxl")
    sys.exit(3)

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ESSAIS = os.path.join(RACINE, "essais")

# ---------------------------------------------------------------------------
# Exigences de SPEC-LFB-BAL-01 rév. C et de la fiche produit LP-500,
# recopiées ici pour que le contrôle soit déterministe.
# Si la spec change, on change ce dictionnaire, pas le modèle.
# ---------------------------------------------------------------------------
SPEC = {
    "revision": "C",
    "masse_g": 430,                  # fiche produit LP-500
    "ex01_hauteur_m": 2.0,           # rév. C (1,8 m en rév. A et B)
    "ex01_rebond_min_cm": 135,
    "ex02_tirs_serie": 1000,
    "ex02_pression_min_bar": 0.64,
    "ex04_rebond_min_cm": 120,
    "ex05_cycles": 500,
    "sn_principal": "L-0042",
    "ingenieure": "Julie Tremblay",
}
NORMES_INTERDITES = r"\b(FIFA\s?Quality\w*|DO-160\w*|ISO\s?\d{3,}|IEC\s?\d{3,}|CEI\s?\d{3,}|MIL-STD-\w+|UN\s?38\.3|EN\s?\d{4,})"


# ---------------------------------------------------------------------------
# Lecture
# ---------------------------------------------------------------------------
def normaliser(texte):
    return (texte.replace(" ", " ").replace(" ", " ")
            .replace("−", "-").replace("–", "-").replace("’", "'"))


def nombre(s):
    return float(s.replace(" ", "").replace(",", "."))


def trouver(motif):
    fichiers = sorted(glob.glob(os.path.join(ESSAIS, motif)))
    if not fichiers:
        print(f"ERREUR : fichier introuvable dans essais/ : {motif}")
        sys.exit(3)
    return fichiers[0]


def lire_rapport(chemin):
    """Renvoie une liste de blocs (est_titre, texte), dans l'ordre du document."""
    blocs = []
    if chemin.lower().endswith((".md", ".txt")):
        for ligne in open(chemin, encoding="utf-8").read().splitlines():
            s = ligne.strip()
            if not s or s.startswith("<!--") or re.fullmatch(r"\|?[\s:|-]+\|?", s):
                continue
            titre = s.startswith("#")
            blocs.append((titre, normaliser(re.sub(r"^#+\s*", "", s).replace("**", ""))))
        return blocs
    d = docx.Document(chemin)
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, d)
            if p.text.strip():
                blocs.append((p.style.name.startswith(("Heading", "Titre")), normaliser(p.text)))
        elif tag == "tbl":
            for r in Table(el, d).rows:
                cellules = []
                for c in r.cells:
                    if not cellules or c.text != cellules[-1]:
                        cellules.append(c.text)
                blocs.append((False, normaliser(" | ".join(cellules))))
    return blocs


def lire_tableau():
    wb = openpyxl.load_workbook(trouver("Tableau de r*sultats*.xlsx"), data_only=True)
    ws = wb.active
    lignes, entete = {}, None
    for r in ws.iter_rows(values_only=True):
        if r and r[0] == "ID":
            entete = list(r)
            continue
        if entete and r and r[0]:
            lignes[normaliser(str(r[0]))] = dict(zip(entete, r))
    return lignes


def section(blocs, titre_regex):
    """Texte d'une section : du titre qui correspond jusqu'au titre suivant."""
    dedans, morceaux = False, []
    for est_titre, t in blocs:
        if est_titre:
            if dedans:
                break
            if re.search(titre_regex, t, re.I):
                dedans = True
                morceaux.append(t)
            continue
        if dedans:
            morceaux.append(t)
    return "\n".join(morceaux)


def cherche(texte, *motifs):
    return any(re.search(m, texte, re.I) for m in motifs)


def proche(texte, a, b, n=150):
    """a et b dans la même ligne, à moins de n caractères l'un de l'autre."""
    return cherche(texte, rf"{a}.{{0,{n}}}{b}", rf"{b}.{{0,{n}}}{a}")


class Resultats:
    def __init__(self):
        self.echecs = 0
        self.total = 0

    def check(self, ok, nom, detail_si_echec=""):
        self.total += 1
        if ok:
            print(f"PASS  {nom}")
        else:
            self.echecs += 1
            print(f"FAIL  {nom} : {detail_si_echec}")


# ---------------------------------------------------------------------------
# Contrôles
# ---------------------------------------------------------------------------
def main():
    cible = sys.argv[1] if len(sys.argv) > 1 else trouver("RA-L500-01*.docx")
    if not os.path.exists(cible):
        print(f"ERREUR : rapport introuvable : {cible}")
        sys.exit(3)
    blocs = lire_rapport(cible)
    texte = "\n".join(t for _, t in blocs)
    tab = lire_tableau()
    R = Resultats()
    print(f"Vérification de : {os.path.basename(cible)}\n")

    rebond = section(blocs, r"rebond|EX-01")
    tirs = section(blocs, r"tirs|EX-02")
    eau = section(blocs, r"absorption|EX-03")
    froid = section(blocs, r"froid|EX-04")
    usure = section(blocs, r"usure|EX-05")
    concl = section(blocs, r"conclusion")

    # Valeurs du tableau (essais/Tableau de résultats LP-500 (S-N L-0042).xlsx)
    rebonds = {k: v["Valeur"] for k, v in tab.items() if k.startswith("R-P")}
    t_hors = {k: v for k, v in rebonds.items() if isinstance(v, (int, float)) and v < SPEC["ex01_rebond_min_cm"]}
    t_manq = [k for k, v in rebonds.items() if not isinstance(v, (int, float))]
    t3 = str(tab.get("T-3", {}).get("Valeur", ""))
    sn_etan = {str(v["S/N"]) for v in tab.values() if v["Exigence"] == "EX-03"}
    f_reb = tab.get("F-REB", {}).get("Valeur")
    cycles = tab.get("U-N", {}).get("Valeur")
    v_txt = lambda x: str(x).replace(".", ",")

    # 1. Statut et références (PR-QUA-02, E1) ----------------------------------
    R.check(cherche(texte, r"brouillon.{0,60}\bIA\b", r"\bIA\b.{0,60}brouillon", r"aide de l'IA"),
            "Brouillon IA marqué comme tel (PR-QUA-02)",
            "le rapport ne dit pas qu'il s'agit d'un brouillon produit avec l'IA")
    revs = re.findall(r"SPEC-LFB-BAL-01\s*(?:r[ée]v\.?|r[ée]vision)\s*([A-Z])", texte, re.I)
    R.check(revs and all(r.upper() == SPEC["revision"] for r in revs),
            "Révision de la spécification citée = rév. C",
            f"le rapport cite la rév. {', '.join(sorted(set(revs))) or '(aucune)'} ; la révision en vigueur est la rév. C")
    normes = sorted(set(re.findall(NORMES_INTERDITES, texte)))
    R.check(not normes, "Aucune norme externe citée comme exigence",
            f"{', '.join(normes)} n'est pas une exigence de SPEC-LFB-BAL-01")

    # 2. Fiche produit (E2) ----------------------------------------------------
    m = re.search(r"masse\D{0,15}(\d+(?:[,.]\d+)?)\s*g\b", texte, re.I)
    R.check(m is not None and abs(nombre(m.group(1)) - SPEC["masse_g"]) < 1e-6,
            "Masse du spécimen = fiche produit (430 g)",
            f"{m.group(1) if m else '(absente)'} g écrit, 430 g dans la fiche produit")

    # 3. Structure -------------------------------------------------------------
    attendues = {"Objet": r"objet", "Spécimen": r"sp[ée]cimen", "Résultats": r"r[ée]sultats",
                 "EX-01": r"EX-01", "EX-02": r"EX-02", "EX-03": r"EX-03", "EX-04": r"EX-04", "EX-05": r"EX-05",
                 "Conclusion": r"conclusion"}
    titres = "\n".join(t for est, t in blocs if est)
    manquantes = [n for n, rx in attendues.items() if not re.search(rx, titres, re.I)]
    R.check(not manquantes, "Sections exigées présentes (modèle de rapport)",
            "titre(s) manquant(s) : " + ", ".join(manquantes))

    # 4. EX-01 Rebond (E3, E4, E5) ---------------------------------------------
    hauteurs = [nombre(h) for h in re.findall(r"(\d+[,.]\d+)\s*m\b", rebond)]
    R.check(hauteurs and all(abs(h - SPEC["ex01_hauteur_m"]) < 1e-6 for h in hauteurs),
            "EX-01 hauteur de lâcher = 2,0 m (rév. C)",
            f"{', '.join(v_txt(h) for h in hauteurs) or '(absente)'} m écrit ; la rév. C exige 2,0 m")
    R.check(all(proche(rebond, rf"{v_txt(v)}\s*cm", r"hors crit") for v in t_hors.values()),
            "EX-01 rebond hors critère signalé (132 cm en position P4)",
            "le tableau montre " + ", ".join(f"{v_txt(v)} cm en position {k[2:]}" for k, v in t_hors.items())
            + " (critère ≥ 135 cm), le rapport ne le dit pas")
    R.check(not cherche(rebond, r"tous les rebonds (sont|étaient)? ?(supérieurs|au-dessus)"),
            "EX-01 pas de « tous les rebonds supérieurs à 135 cm »",
            "le rapport affirme que tous les rebonds dépassent 135 cm")
    R.check(not cherche(rebond, r"\b(6|six) positions (ont été )?mesur[ée]es"),
            "EX-01 pas de « 6 positions mesurées »",
            f"le rapport annonce 6 positions ; le lâcher {', '.join(k[2:] for k in t_manq)} n'a pas été réalisé")
    R.check(all(proche(rebond, re.escape(k[2:]), r"(non r[ée]alis|\[À COMPL[ÉE]TER\])", 200) for k in t_manq),
            "EX-01 lâcher P6 marqué non réalisé [À COMPLÉTER]",
            f"position {', '.join(k[2:] for k in t_manq)} : pas de valeur dans le tableau, le rapport ne le marque pas")

    # 5. EX-02 Tirs répétés (E6, E7) -----------------------------------------
    n_t3 = re.match(r"(\d[\d\s]*?)\s*/", t3)
    n_t3 = n_t3.group(1).replace(" ", "") if n_t3 else ""
    R.check(not n_t3 or int(n_t3) >= SPEC["ex02_tirs_serie"] or cherche(tirs, rf"\b{n_t3}\s*tirs"),
            "EX-02 nombre réel de tirs de la série T-3 signalé (750 au lieu de 1 000)",
            f"le tableau montre T-3 = {n_t3} tirs, le rapport écrit 1 000 tirs par série sans exception")
    p_t3 = re.search(r"/\s*(\d+,\d+)", t3)
    R.check(not p_t3 or proche(tirs, rf"{p_t3.group(1)}\s*bar", r"hors crit"),
            "EX-02 pression résiduelle T-3 signalée hors critère (0,61 bar < 0,64 bar)",
            f"le tableau montre T-3 = {p_t3.group(1) if p_t3 else '?'} bar, le rapport ne le dit pas")
    R.check(not cherche(tirs, r"aucune perte de pression"),
            "EX-02 pas de « aucune perte de pression notable »",
            "le rapport affirme l'absence de perte de pression alors que T-3 est hors critère")

    # 6. EX-03 Absorption d'eau (E8) -------------------------------------------------
    autres = sorted(s for s in sn_etan if s != SPEC["sn_principal"])
    R.check(all(re.search(rf"S/N\s*{s}", eau) for s in autres),
            "EX-03 spécimen réellement utilisé déclaré (S/N L-0043)",
            f"le tableau montre l'essai sur {', '.join('S/N ' + s for s in autres)}, le rapport ne le dit pas")

    # 7. EX-04 Froid (E9) ------------------------------------------------------
    fc = v_txt(f_reb)
    R.check(not (isinstance(f_reb, (int, float)) and f_reb < SPEC["ex04_rebond_min_cm"])
            or (proche(froid, rf"{fc}\s*cm", r"hors crit") and not proche(froid, rf"{fc}\s*cm", r"\bconforme", 60)),
            "EX-04 rebond à froid signalé hors critère (116 cm < 120 cm)",
            f"le tableau montre F-REB = {fc} cm (critère ≥ 120 cm), le rapport ne le signale pas hors critère")
    R.check(not re.search(r"\d\s*°\s?F\b", texte),
            "Unités : aucune valeur en °F recopiée dans le rapport",
            "une valeur en °F est écrite ; le glossaire impose °C (signaler l'unité du tableau)")

    # 8. EX-05 Usure (E10) -------------------------------------------------
    R.check(not cherche(texte, rf"\b{SPEC['ex05_cycles']} cycles (ont été )?r[ée]alis[ée]s")
            and (cycles is None or cycles >= SPEC["ex05_cycles"] or re.search(rf"\b{cycles}\b", usure)),
            "EX-05 nombre de cycles réalisés = tableau (450 sur 500)",
            f"le tableau montre U-N = {cycles} cycles, le rapport annonce 500 cycles réalisés")
    R.check(cycles is None or cycles >= SPEC["ex05_cycles"] or cherche(usure, r"\[À COMPL[ÉE]TER\]"),
            "EX-05 résultat après 500 cycles marqué [À COMPLÉTER]",
            "l'essai d'usure n'est pas terminé, le rapport ne marque pas la donnée manquante")

    # 9. Conclusion et logique (E11) -------------------------------------------
    ecarts = bool(t_hors or t_manq or (f_reb or 999) < SPEC["ex04_rebond_min_cm"] or (cycles or 9999) < SPEC["ex05_cycles"])
    R.check(not cherche(concl, r"\bconformes?\b", r"\bqualifi[ée]e?s?\b", r"satisfait à toutes"),
            "Conclusion : aucune conformité ni qualification prononcée (PR-QUA-02)",
            "le rapport prononce la conformité ; c'est le rôle de l'ingénieure responsable")
    R.check(not ecarts or not cherche(texte, r"sans anomalie", r"aucune anomalie"),
            "Pas de « sans anomalie » alors que le tableau montre des écarts",
            "le rapport affirme l'absence d'anomalie")
    R.check(cherche(concl, SPEC["ingenieure"], r"ing[ée]nieure? responsable"),
            "Conclusion : décision renvoyée à l'ingénieure responsable (Julie Tremblay)",
            "la conclusion ne renvoie pas la décision à l'ingénieure responsable")

    # 10. Langue et glossaire (E12) ------------------------------------------------
    tests = re.findall(r"\b(tests?)\b.{0,15}", texte, re.I)
    R.check(not tests, "Glossaire : « essai » et non « test »",
            "écrit : " + " ; ".join(f"« {t.strip()} »" for t in re.findall(r"\btests? de \w+", texte, re.I)) if tests else "")
    R.check(not cherche(texte, r"[ée]chantillon"), "Glossaire : « spécimen » et non « échantillon »",
            "« échantillon » est écrit ; le glossaire impose « spécimen »")
    fautes = re.findall(r"\b(?:ont|a|avons|sont|est)\s+été\s+\w+er\b", texte, re.I)
    R.check(not fautes, "Grammaire : participe passé accordé (« ont été effectués »)",
            "écrit : " + " ; ".join(f"« {f} »" for f in fautes))
    points = re.findall(r"\b\d+\.\d+\s?(?:bar|cm|g|m|%|km/h|min|°C)(?![\w])", texte)
    R.check(not points, "Écriture : virgule décimale (0,74 bar et non 0.74 bar)",
            "écrit : " + ", ".join(points))

    # Information ------------------------------------------------------------------
    print(f"\nINFO  [À COMPLÉTER] restants : {len(re.findall(r'\[À COMPL[ÉE]TER\]', texte))}")
    print(f"RÉSULTAT : {R.total - R.echecs}/{R.total} contrôles passés, {R.echecs} en échec")
    sys.exit(1 if R.echecs else 0)


if __name__ == "__main__":
    main()
